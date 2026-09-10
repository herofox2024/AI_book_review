"""Word (.docx) 解析：按标题样式识别章节。"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from .base import BaseParser
from ..models import Book


class DocxParser(BaseParser):
    formats = (".docx",)

    def parse(self, path: Path) -> Book:
        try:
            doc = Document(str(path))
        except Exception as e:
            raise RuntimeError(f"无法读取 docx（若文件损坏请另存为新文件）: {e}") from e

        title = (doc.core_properties.title or "").strip()
        author = (doc.core_properties.author or "").strip()

        blocks: list[tuple[str, str]] = []
        cur_title = ""
        cur_buf: list[str] = []

        def flush() -> None:
            if cur_buf:
                blocks.append((cur_title, "\n".join(cur_buf)))
                cur_buf.clear()

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            is_heading = style.startswith("heading") or style.startswith("标题")
            if is_heading and len(text) <= 60:
                flush()
                cur_title = text
            else:
                cur_buf.append(text)
        flush()

        # 表格中的文字也纳入，避免丢失正文
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    cur_buf.append(" ".join(cells))
        if cur_buf:
            blocks.append((cur_title or "表格内容", "\n".join(cur_buf)))

        if not blocks:
            raise RuntimeError("docx 中未找到任何文本")

        return self.build_book(path, title, author, "docx", blocks, language="zh")
