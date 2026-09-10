"""生成测试用电子书样本：docx / epub。"""

from pathlib import Path

from docx import Document
from ebooklib import epub

OUT = Path(__file__).parent / "sample_books"
OUT.mkdir(exist_ok=True)

CHAPTERS = [
    ("第一章 雪夜来客", [
        "那年冬天的雪下得格外早，十一月中旬，东京就已经白茫茫一片了。",
        "光司先生推开门的时候，肩上还落着未化的雪。他抖了抖大衣，动作很轻，仿佛怕惊扰了屋里的什么。",
        "「打扰了。」他说。声音比记忆中低沉了些。",
        "我盯着他看了很久。七年了。七年前那个在法庭上一言不发的青年，如今眉眼间多了一种我说不清的东西。",
        "「你还愿意接这个案子吗？」他终于抬起头。",
        "我没有立刻回答。窗外的雪落在玻璃上，化成一道道水痕，像谁在无声地写字。",
    ]),
    ("第二章 薄冰之沼", [
        "案子发生在沼津。一片被当地人称作「薄冰之沼」的浅水塘。",
        "尸体是十二月三日清晨被钓鱼的老人发现的。女性，二十七岁，名叫佐藤美咲。",
        "光司先生是她的未婚夫。这一点，他第二天才告诉我。",
        "「我以为不说，就不会被怀疑。」他说这话时手指在膝上收紧，指节发白。",
        "我合上卷宗。第一页的照片拍得很清楚——死者的指甲缝里，有少量蓝色的纤维。",
        "而光司先生的围巾，正是深蓝色的。",
    ]),
    ("第三章 证词", [
        "第一个证人是死者的同事，一个说话很快的年轻女人。",
        "「美咲最近很怕。」她说，「十一月底开始，她总说有人跟着她。」",
        "「谁？」我问。",
        "「不知道。她不肯说。但她把婚礼取消了。」",
        "这句话让整个房间安静下来。光司先生坐在角落里，一动不动。",
        "我忽然意识到，我们一直在问错问题。重要的不是谁杀了她，而是她在怕什么。",
    ]),
    ("第四章 第七年的雪", [
        "真相揭晓的那天，雪又下了。",
        "跟踪美咲的人，是她自己的哥哥。理由荒唐得令人发笑——他不同意这门婚事。",
        "而杀死她的，却是另一个人。一个我们从头到尾都忽略的人。",
        "光司先生在警局里坐了整整一夜。他没有哭，只是反复说同一句话：「我本该早一点告诉她。」",
        "我走出警局时，雪已经停了。薄冰之沼上结了一层新的冰，看上去很结实，其实一踩就碎。",
        "人和人之间，大抵也是如此。",
    ]),
]


def make_docx() -> Path:
    doc = Document()
    doc.core_properties.title = "薄冰之沼"
    doc.core_properties.author = "笹泽左保"
    doc.add_heading("薄冰之沼", level=0)
    for title, paras in CHAPTERS:
        doc.add_heading(title, level=1)
        for p in paras:
            doc.add_paragraph(p)
    path = OUT / "薄冰之沼.docx"
    doc.save(str(path))
    return path


def make_epub() -> Path:
    book = epub.EpubBook()
    book.set_identifier("test-thin-ice-001")
    book.set_title("薄冰之沼")
    book.set_language("zh")
    book.add_author("笹泽左保")

    spine = ["nav"]
    toc = []
    for i, (title, paras) in enumerate(CHAPTERS):
        html = f"<h1>{title}</h1>" + "".join(f"<p>{p}</p>" for p in paras)
        ch = epub.EpubHtml(title=title, file_name=f"chap_{i+1}.xhtml", lang="zh")
        ch.content = html
        book.add_item(ch)
        spine.append(ch)
        toc.append(ch)

    book.toc = tuple(toc)
    book.spine = spine
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    path = OUT / "薄冰之沼.epub"
    epub.write_epub(str(path), book)
    return path


if __name__ == "__main__":
    print("docx:", make_docx())
    print("epub:", make_epub())
